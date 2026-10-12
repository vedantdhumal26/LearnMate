"""
Complete Phase 4A Live Acceptance Test for LearnSphere AI.
Multimodal Ingestion: PPTX + Slide-Level Citations + Visual/Diagram Extraction.
Executes all 11 steps against the live running FastAPI server at http://127.0.0.1:8000.
"""

import os
import sys
import io

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import httpx
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

BASE_URL = "http://127.0.0.1:8000"

def generate_acceptance_pptx() -> bytes:
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title and Overview
    slide1 = prs.slides.add_slide(blank_layout)
    box1 = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
    tf1 = box1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Introduction to Aerodynamics and Fluid Flow"
    p1.font.size = Pt(28)
    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Fundamental principles of fluid dynamics, pressure differentials, and laminar flow."

    # Slide 2: Unique Bernoulli Content (Answer to question exists ONLY here)
    slide2 = prs.slides.add_slide(blank_layout)
    box2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "Bernoulli's Principle and the Venturi Tube Effect"
    p2.font.size = Pt(24)
    p2_content = tf2.add_paragraph()
    p2_content.text = (
        "According to Bernoulli's principle, as the velocity of a fluid increases, the static pressure "
        "exerted by the fluid decreases simultaneously. The governing equation is P + 0.5 * rho * v^2 + rho * g * h = constant. "
        "Venturi tubes utilize this exact velocity increase and pressure drop for fluid metering and fuel atomization in carburetors."
    )
    p2_formula = tf2.add_paragraph()
    p2_formula.text = "Static Pressure decreases as Velocity increases along a streamline."

    # Slide 3: Embedded Visual / Diagram
    slide3 = prs.slides.add_slide(blank_layout)
    box3 = slide3.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tf3 = box3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = "Airfoil Lift and Streamline Pressure Distribution"

    # Create synthetic diagram image
    img = Image.new("RGB", (300, 200), color=(40, 100, 180))
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    img_byte_arr.seek(0)
    slide3.shapes.add_picture(img_byte_arr, Inches(1), Inches(1.5), Inches(5), Inches(3.5))

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()

def run_acceptance_test():
    print("=" * 80)
    print("STARTING PHASE 4A LIVE ACCEPTANCE TEST: MULTIMODAL PPTX INGESTION & CITATIONS")
    print("=" * 80)

    client = httpx.Client(base_url=BASE_URL, timeout=60.0)

    # 1. Create a course
    print("\n[Step 1] Creating course...")
    c1_res = client.post("/api/courses", json={
        "title": "Aerodynamics and Fluid Dynamics",
        "description": "Study of fluid flow, lift, drag, and Bernoulli principles."
    })
    assert c1_res.status_code == 201, f"Failed to create course: {c1_res.text}"
    course_1 = c1_res.json()
    course_1_id = course_1["id"]
    print(f"--> Course 1 created: ID={course_1_id}, Title='{course_1['title']}'")

    # 2. Generate a real PPTX
    print("\n[Step 2] Generating real PPTX with 3 slides (text, Bernoulli formula, embedded diagram)...")
    pptx_bytes = generate_acceptance_pptx()
    print(f"--> Generated PPTX package size: {len(pptx_bytes)} bytes")

    # 3. Upload PPTX
    print("\n[Step 3] Uploading PPTX to Course 1...")
    files = {
        "file": (
            "aerodynamics_lecture.pptx",
            io.BytesIO(pptx_bytes),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
    }
    upload_res = client.post(f"/api/courses/{course_1_id}/upload", files=files)
    assert upload_res.status_code == 201, f"Failed to upload PPTX: {upload_res.text}"
    upload_data = upload_res.json()
    source_id = upload_data["source_id"]
    file_name = upload_data["file_name"]
    print(f"--> Upload successful: source_id={source_id}, file_name='{file_name}', status='{upload_data.get('status')}'")

    # 4. Verify processing_status = completed
    print("\n[Step 4] Checking ingestion status...")
    status_res = client.get(f"/api/courses/{course_1_id}/uploads/{source_id}")
    assert status_res.status_code == 200, f"Failed to get source status: {status_res.text}"
    source_info = status_res.json()
    assert source_info["processing_status"] == "completed", f"Status is {source_info['processing_status']}"
    print(f"--> Source processing_status: {source_info['processing_status']}")

    # 5. Verify slide count / chunks
    print("\n[Step 5] Verifying chunks and slide metadata...")
    chunks_count = source_info.get("chunks_count", 0)
    slide_count = source_info.get("slide_count")
    visual_count = source_info.get("visual_count")
    print(f"--> Chunks count: {chunks_count}, Slide count: {slide_count}, Visual count: {visual_count}")
    assert chunks_count >= 3, f"Expected at least 3 chunks, got {chunks_count}"
    assert slide_count == 3, f"Expected 3 slides, got {slide_count}"

    # 6. Ask a question whose answer exists only on Slide 2
    print("\n[Step 6] Asking question answered ONLY on Slide 2...")
    q1 = "According to Bernoulli's principle, what happens to the static pressure of a fluid when its velocity increases?"
    chat_res = client.post("/api/tutor/chat", json={
        "course_id": course_1_id,
        "message": q1
    })
    assert chat_res.status_code == 200, f"Tutor chat failed: {chat_res.text}"
    chat_data = chat_res.json()

    # 7. Verify grounded = true, source_id correct, file_name correct, slide = 2
    print("\n[Step 7] Verifying tutor answer and slide-level citation...")
    print(f"--> Grounded: {chat_data['grounded']}")
    print(f"--> Answer: {chat_data['answer']}")
    assert chat_data["grounded"] is True, "Expected tutor answer to be grounded"
    assert len(chat_data["sources"]) > 0, "Expected at least one citation source"

    slide_citation = None
    for src in chat_data["sources"]:
        print(f"--> Citation: file='{src['file_name']}', slide={src.get('slide')}, page={src.get('page')}")
        if src.get("slide") == 2:
            slide_citation = src
            break

    assert slide_citation is not None, f"Expected citation with slide=2, got citations: {chat_data['sources']}"
    assert slide_citation["source_id"] == source_id, f"Expected source_id={source_id}, got {slide_citation['source_id']}"
    assert slide_citation["file_name"] == "aerodynamics_lecture.pptx"
    assert slide_citation["page"] is None, f"Expected page=None for PPTX, got {slide_citation['page']}"
    assert slide_citation["slide"] == 2, f"Expected slide=2, got {slide_citation['slide']}"
    print("--> PASS: Verified exact slide citation: aerodynamics_lecture.pptx — Slide 2")

    # 8. Ask an unrelated question
    print("\n[Step 8] Asking completely unrelated off-material question...")
    q_unrelated = "What is the capital of ancient Mesopotamia and what were the laws of Hammurabi?"
    chat_unrelated = client.post("/api/tutor/chat", json={
        "course_id": course_1_id,
        "message": q_unrelated
    })
    assert chat_unrelated.status_code == 200
    unrelated_data = chat_unrelated.json()

    # 9. Verify grounded = false, sources = []
    print("\n[Step 9] Verifying refusal for off-material question...")
    print(f"--> Grounded: {unrelated_data['grounded']}")
    print(f"--> Sources: {unrelated_data['sources']}")
    print(f"--> Answer: {unrelated_data['answer']}")
    assert unrelated_data["grounded"] is False, "Expected grounded=False for off-material topic"
    assert unrelated_data["sources"] == [], "Expected empty sources for off-material topic"
    print("--> PASS: Off-material question correctly rejected with grounded=False and sources=[]")

    # 10. Create second course
    print("\n[Step 10] Creating second course to test cross-course isolation...")
    c2_res = client.post("/api/courses", json={
        "title": "Astrophysics and Stellar Evolution",
        "description": "Lifecycle of stars, white dwarfs, neutron stars, and black holes."
    })
    assert c2_res.status_code == 201
    course_2 = c2_res.json()
    course_2_id = course_2["id"]
    print(f"--> Course 2 created: ID={course_2_id}, Title='{course_2['title']}'")

    # 11. Verify Course 1 content cannot be retrieved from Course 2
    print("\n[Step 11] Verifying Course 1 content is completely inaccessible from Course 2...")
    cross_res = client.post("/api/tutor/chat", json={
        "course_id": course_2_id,
        "message": "Explain Bernoulli's principle and the governing pressure equation."
    })
    assert cross_res.status_code == 200
    cross_data = cross_res.json()
    print(f"--> Course 2 Query Grounded: {cross_data['grounded']}")
    print(f"--> Course 2 Query Sources: {cross_data['sources']}")
    assert cross_data["grounded"] is False, "Course 1 content must NOT be retrievable from Course 2"
    assert cross_data["sources"] == [], "No sources should be returned across course boundary"
    print("--> PASS: Strict Course Isolation verified. Course 2 returned grounded=False.")

    print("\n" + "=" * 80)
    print("ALL 11 STEPS OF LIVE ACCEPTANCE TEST COMPLETED SUCCESSFULLY!")
    print("=" * 80)

if __name__ == "__main__":
    run_acceptance_test()
