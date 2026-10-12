import io
import pytest
from unittest.mock import patch, MagicMock
from app.models.course import Course
from app.models.source import Source
from app.models.document_chunk import DocumentChunk
from app.services.transcription.transcription_service import (
    TranscriptionService,
    TranscriptionResult,
    TranscriptionSegment,
    format_timestamp
)
from app.services.vector_store.qdrant_store import get_vector_store
from tests.conftest import MockLLMService

def generate_synthetic_mp4() -> bytes:
    """Generate a minimal valid MP4 container binary payload with standard ftyp box."""
    ftyp_data = b"mp42\x00\x00\x00\x00isommp42"
    ftyp_box = (len(ftyp_data) + 8).to_bytes(4, "big") + b"ftyp" + ftyp_data
    payload = b"synthetic video and audio frame sample payload"
    mdat_box = (len(payload) + 8).to_bytes(4, "big") + b"mdat" + payload
    return ftyp_box + mdat_box

@pytest.fixture
def sample_mp4() -> bytes:
    return generate_synthetic_mp4()

@pytest.fixture
def mock_transcription_result() -> TranscriptionResult:
    return TranscriptionResult(
        segments=[
            TranscriptionSegment(
                text="Welcome to fluid mechanics. Today we examine fluid pressure and velocity.",
                start=0.0,
                end=15.2
            ),
            TranscriptionSegment(
                text="Bernoulli's principle states that as fluid velocity increases, static pressure decreases simultaneously.",
                start=15.2,
                end=42.5
            ),
            TranscriptionSegment(
                text="The governing relation is P plus half rho v squared plus rho g h equals a constant along streamlines.",
                start=42.5,
                end=85.0
            ),
            TranscriptionSegment(
                text="Venturi meters demonstrate this pressure drop when a fluid constricts through a throat section.",
                start=85.0,
                end=120.0
            )
        ],
        duration_seconds=120.0,
        full_text="Welcome to fluid mechanics. Today we examine fluid pressure and velocity. Bernoulli's principle states that as fluid velocity increases, static pressure decreases simultaneously. The governing relation is P plus half rho v squared plus rho g h equals a constant along streamlines. Venturi meters demonstrate this pressure drop when a fluid constricts through a throat section."
    )

def test_mp4_upload_validation(client):
    """1. Test validation on empty and corrupted MP4 files."""
    course_res = client.post("/api/courses", json={"title": "Validation Course"})
    course_id = course_res.json()["id"]

    # Empty file
    empty_file = {"file": ("empty.mp4", io.BytesIO(b""), "video/mp4")}
    res_empty = client.post(f"/api/courses/{course_id}/upload", files=empty_file)
    assert res_empty.status_code == 400
    assert "empty" in res_empty.json()["detail"]

    # Corrupt container header
    corrupt_file = {"file": ("corrupt.mp4", io.BytesIO(b"not a valid container"), "video/mp4")}
    res_corrupt = client.post(f"/api/courses/{course_id}/upload", files=corrupt_file)
    assert res_corrupt.status_code == 400
    assert "Invalid MP4 structure" in res_corrupt.json()["detail"]

def test_mp4_upload_success(client, sample_mp4, mock_transcription_result):
    """2. Test successful MP4 upload and processing status response."""
    course_res = client.post("/api/courses", json={"title": "Fluid Dynamics Video"})
    course_id = course_res.json()["id"]

    with patch.object(TranscriptionService, "transcribe_file", return_value=mock_transcription_result):
        files = {"file": ("lecture1.mp4", io.BytesIO(sample_mp4), "video/mp4")}
        res = client.post(f"/api/courses/{course_id}/upload", files=files)
        assert res.status_code == 201
        data = res.json()
        assert data["file_name"] == "lecture1.mp4"
        assert data["status"] == "completed"

        # Check status endpoint
        status_res = client.get(f"/api/courses/{course_id}/uploads/{data['source_id']}")
        assert status_res.status_code == 200
        status_data = status_res.json()
        assert status_data["processing_status"] == "completed"
        assert status_data["duration_seconds"] is not None
        assert status_data["duration_seconds"] > 0
        assert status_data["segment_count"] is not None
        assert status_data["segment_count"] >= 1

def test_transcription_service_response_normalization():
    """3. Test normalization of raw segments including malformed inputs."""
    service = TranscriptionService(api_key="mock-key")
    raw = [
        {"text": "  Valid text segment.  ", "start": "10.5", "end": "22.3"},
        {"text": "", "start": 23.0, "end": 25.0}, # Empty text should be skipped
        {"text": "Negative or reversed duration", "start": 30.0, "end": 28.0}, # end adjusted
        {"bad_key": 123}, # Malformed segment skipped
        MagicMock(text="Object segment", start=40.0, end=50.0) # Object attribute access
    ]
    normalized = service.normalize_segments(raw)
    assert len(normalized) == 3
    assert normalized[0].text == "Valid text segment."
    assert normalized[0].start == 10.5
    assert normalized[0].end == 22.3
    assert normalized[1].text == "Negative or reversed duration"
    assert normalized[1].end >= normalized[1].start
    assert normalized[2].text == "Object segment"

def test_video_segment_timestamps_preserved(client, db_session, sample_mp4, mock_transcription_result):
    """4. Test that start_time_seconds and end_time_seconds are preserved on chunks and page/slide are None."""
    course_res = client.post("/api/courses", json={"title": "Timestamp Preserved Course"})
    course_id = course_res.json()["id"]

    with patch.object(TranscriptionService, "transcribe_file", return_value=mock_transcription_result):
        files = {"file": ("timestamps.mp4", io.BytesIO(sample_mp4), "video/mp4")}
        res = client.post(f"/api/courses/{course_id}/upload", files=files)
        source_id = res.json()["source_id"]

        chunks = db_session.query(DocumentChunk).filter(DocumentChunk.source_id == source_id).all()
        assert len(chunks) >= 1
        for chunk in chunks:
            assert chunk.start_time_seconds is not None
            assert chunk.end_time_seconds is not None
            assert chunk.end_time_seconds > chunk.start_time_seconds
            assert chunk.page_number is None
            assert chunk.slide_number is None
            assert chunk.modality == "video_transcript"

def test_video_chunks_created(client, db_session, sample_mp4, mock_transcription_result):
    """5. Test that video chunks have formatted video headers, content hashes, and correct bounds."""
    course_res = client.post("/api/courses", json={"title": "Chunks Created Course"})
    course_id = course_res.json()["id"]

    with patch.object(TranscriptionService, "transcribe_file", return_value=mock_transcription_result):
        files = {"file": ("chunks.mp4", io.BytesIO(sample_mp4), "video/mp4")}
        res = client.post(f"/api/courses/{course_id}/upload", files=files)
        source_id = res.json()["source_id"]

        chunks = db_session.query(DocumentChunk).filter(DocumentChunk.source_id == source_id).all()
        for chunk in chunks:
            assert "[Video " in chunk.content
            assert len(chunk.content_hash) == 64
            assert chunk.course_id == course_id

def test_video_vectors_indexed(client, sample_mp4, mock_transcription_result):
    """6. Test that video vectors are indexed in Qdrant with timestamps and modality."""
    course_res = client.post("/api/courses", json={"title": "Vector Video Course"})
    course_id = course_res.json()["id"]

    with patch.object(TranscriptionService, "transcribe_file", return_value=mock_transcription_result):
        files = {"file": ("vectors.mp4", io.BytesIO(sample_mp4), "video/mp4")}
        client.post(f"/api/courses/{course_id}/upload", files=files)

        store = get_vector_store()
        results = store.search_vectors(
            query_vector=[0.0] * 384,
            course_id=course_id,
            limit=10,
            score_threshold=0.0
        )
        assert len(results) >= 1
        for r in results:
            assert r.file_name == "vectors.mp4"
            assert r.modality == "video_transcript"
            assert r.start_time_seconds is not None
            assert r.end_time_seconds is not None

def test_video_tutor_grounded_answer(client, sample_mp4, mock_transcription_result):
    """7. Test grounded tutor answer generated from video transcript content."""
    course_res = client.post("/api/courses", json={"title": "Tutor Video Course"})
    course_id = course_res.json()["id"]

    with patch.object(TranscriptionService, "transcribe_file", return_value=mock_transcription_result):
        files = {"file": ("lecture_fluid.mp4", io.BytesIO(sample_mp4), "video/mp4")}
        client.post(f"/api/courses/{course_id}/upload", files=files)

        mock_llm = MockLLMService(response_text="According to lecture_fluid.mp4 — 00:00:15–00:00:42, Bernoulli's principle states that as fluid velocity increases, static pressure decreases.")
        with patch("app.api.routes.tutor.tutor_service._llm_service", mock_llm):
            chat_res = client.post("/api/tutor/chat", json={
                "course_id": course_id,
                "message": "What does Bernoulli's principle state in this lecture?"
            })
            assert chat_res.status_code == 200
            data = chat_res.json()
            assert data["grounded"] is True
            assert "Bernoulli" in data["answer"]
            assert len(data["sources"]) > 0

def test_video_tutor_returns_timestamp_citation(client, sample_mp4, mock_transcription_result):
    """8. Test that tutor returns exact timestamp citation (start_seconds, end_seconds, slide=null, page=null)."""
    course_res = client.post("/api/courses", json={"title": "Citation Video Course"})
    course_id = course_res.json()["id"]

    with patch.object(TranscriptionService, "transcribe_file", return_value=mock_transcription_result):
        files = {"file": ("lecture1.mp4", io.BytesIO(sample_mp4), "video/mp4")}
        client.post(f"/api/courses/{course_id}/upload", files=files)

        mock_llm = MockLLMService(response_text="The governing relation is P + 0.5 rho v^2 = constant as stated in lecture1.mp4 — 00:00:42–00:01:25.")
        with patch("app.api.routes.tutor.tutor_service._llm_service", mock_llm):
            chat_res = client.post("/api/tutor/chat", json={
                "course_id": course_id,
                "message": "governing equation P plus half rho v squared"
            })
            assert chat_res.status_code == 200
            data = chat_res.json()
            assert data["grounded"] is True
            assert len(data["sources"]) > 0
            citation = data["sources"][0]
            assert citation["file_name"] == "lecture1.mp4"
            assert citation["page"] is None
            assert citation["slide"] is None
            assert citation["start_seconds"] is not None
            assert citation["end_seconds"] is not None
            assert citation["end_seconds"] > citation["start_seconds"]

def test_video_course_isolation(client, sample_mp4, mock_transcription_result):
    """9. Test strict course isolation with video materials."""
    c1_res = client.post("/api/courses", json={"title": "Video Course 1"})
    c1_id = c1_res.json()["id"]
    c2_res = client.post("/api/courses", json={"title": "Video Course 2 Empty"})
    c2_id = c2_res.json()["id"]

    with patch.object(TranscriptionService, "transcribe_file", return_value=mock_transcription_result):
        files = {"file": ("secret_fluid.mp4", io.BytesIO(sample_mp4), "video/mp4")}
        client.post(f"/api/courses/{c1_id}/upload", files=files)

        # Ask in empty Course 2
        chat_res = client.post("/api/tutor/chat", json={
            "course_id": c2_id,
            "message": "What does Bernoulli's principle state in the lecture?"
        })
        assert chat_res.status_code == 200
        data = chat_res.json()
        assert data["grounded"] is False
        assert data["sources"] == []

def test_pdf_regression(client, sample_physics_pdf):
    """10. Test PDF ingestion and citation continues to work normally."""
    course_res = client.post("/api/courses", json={"title": "PDF Regression Test"})
    course_id = course_res.json()["id"]

    files = {"file": ("physics.pdf", io.BytesIO(sample_physics_pdf), "application/pdf")}
    res = client.post(f"/api/courses/{course_id}/upload", files=files)
    assert res.status_code == 201

    mock_llm = MockLLMService(response_text="Newton's second law is F = ma on physics.pdf — Page 2.")
    with patch("app.api.routes.tutor.tutor_service._llm_service", mock_llm):
        chat_res = client.post("/api/tutor/chat", json={
            "course_id": course_id,
            "message": "F = m * a force equation"
        })
        assert chat_res.status_code == 200
        data = chat_res.json()
        assert data["grounded"] is True
        assert len(data["sources"]) > 0
        citation = data["sources"][0]
        assert citation["page"] is not None
        assert citation["slide"] is None
        assert citation["start_seconds"] is None

def test_pptx_regression(client):
    """11. Test PPTX ingestion and slide citation continues to work normally."""
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tx = slide.shapes.add_textbox(0, 0, 100, 100)
    tx.text_frame.text = "Thermodynamics Lecture on Kinetic Gas Theory"
    buf = io.BytesIO()
    prs.save(buf)

    course_res = client.post("/api/courses", json={"title": "PPTX Regression Test"})
    course_id = course_res.json()["id"]

    files = {"file": ("slides.pptx", io.BytesIO(buf.getvalue()), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res = client.post(f"/api/courses/{course_id}/upload", files=files)
    assert res.status_code == 201

    mock_llm = MockLLMService(response_text="Covered on slides.pptx — Slide 1.")
    with patch("app.api.routes.tutor.tutor_service._llm_service", mock_llm):
        chat_res = client.post("/api/tutor/chat", json={
            "course_id": course_id,
            "message": "Thermodynamics kinetic gas theory"
        })
        assert chat_res.status_code == 200
        data = chat_res.json()
        assert data["grounded"] is True
        assert len(data["sources"]) > 0
        citation = data["sources"][0]
        assert citation["slide"] == 1
        assert citation["page"] is None
        assert citation["start_seconds"] is None

def test_transcription_failure_marks_source_failed(client, db_session, sample_mp4):
    """12. Test that transcription failure marks source processing_status = 'failed' and returns 500."""
    course_res = client.post("/api/courses", json={"title": "Failure Test Course"})
    course_id = course_res.json()["id"]

    with patch.object(TranscriptionService, "transcribe_file", side_effect=RuntimeError("Whisper quota exceeded")):
        files = {"file": ("failing_video.mp4", io.BytesIO(sample_mp4), "video/mp4")}
        res = client.post(f"/api/courses/{course_id}/upload", files=files)
        assert res.status_code == 500

        source = db_session.query(Source).filter(Source.course_id == course_id, Source.file_name == "failing_video.mp4").first()
        assert source is not None
        assert source.processing_status == "failed"

def test_timestamp_offset_for_chunked_audio():
    """13. Test that chunking audio accurately offsets segment timestamps by audio chunk start."""
    service = TranscriptionService(api_key="mock-key")
    raw = [
        {"text": "Segment 1 in chunk 2", "start": 2.5, "end": 7.0},
        {"text": "Segment 2 in chunk 2", "start": 8.0, "end": 14.5}
    ]
    # Chunk 2 begins at 600.0s
    offset_segments = service.normalize_segments(raw, offset_seconds=600.0)
    assert len(offset_segments) == 2
    assert offset_segments[0].start == 602.5
    assert offset_segments[0].end == 607.0
    assert offset_segments[1].start == 608.0
    assert offset_segments[1].end == 614.5

def test_unsupported_video_extension_rejected(client):
    """14. Test that unsupported video extensions like .avi or .mkv are rejected with Phase 4B message."""
    course_res = client.post("/api/courses", json={"title": "Unsupported Video Course"})
    course_id = course_res.json()["id"]

    files_avi = {"file": ("lecture.avi", io.BytesIO(b"RIFF....AVI "), "video/x-msvideo")}
    res_avi = client.post(f"/api/courses/{course_id}/upload", files=files_avi)
    assert res_avi.status_code == 400
    assert res_avi.json()["detail"] == "Only .mp4 video files are supported in Phase 4B."

    files_mkv = {"file": ("lecture.mkv", io.BytesIO(b"\x1a\x45\xdf\xa3"), "video/x-matroska")}
    res_mkv = client.post(f"/api/courses/{course_id}/upload", files=files_mkv)
    assert res_mkv.status_code == 400
    assert res_mkv.json()["detail"] == "Only .mp4 video files are supported in Phase 4B."
