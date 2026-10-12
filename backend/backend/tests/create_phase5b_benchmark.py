import os
import sys
import time
import json
import subprocess
from typing import Dict, Any

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import httpx
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt
from app.services.transcription.audio_utils import get_ffmpeg_exe

BASE_URL = "http://127.0.0.1:8000"
REAL_YOUTUBE_URL = "https://www.youtube.com/watch?v=aircAruvnKk"


def create_10_page_pdf(file_path: str) -> str:
    """Create a 10-page comprehensive Physics Guide PDF."""
    doc = pymupdf.open()
    pages_content = [
        ("Chapter 1: Newton's First Law of Motion",
         "Newton's First Law of Motion states that an object at rest stays at rest and an object in motion "
         "continues in uniform motion along a straight line unless acted upon by an external net force. "
         "This property of matter resisting changes in its velocity is known as inertia. Mass is the quantitative measure of inertia."),
        
        ("Chapter 2: Newton's Second Law of Motion",
         "Newton's Second Law of Motion states that the acceleration of an object is directly proportional "
         "to the net force acting upon it and inversely proportional to its mass. "
         "The fundamental mathematical equation is F = m * a, where F is net force in Newtons, m is mass in kilograms, "
         "and a is acceleration in meters per second squared."),
        
        ("Chapter 3: Newton's Third Law of Motion",
         "Newton's Third Law of Motion states that when one body exerts a force on a second body, "
         "the second body simultaneously exerts a force equal in magnitude and opposite in direction on the first body. "
         "In short: to every action there is always an equal and opposite reaction. Action-reaction pairs never act on the same object."),
        
        ("Chapter 4: Conservation of Linear Momentum",
         "Linear momentum is defined as the product of an object's mass and its velocity: p = m * v. "
         "The principle of conservation of linear momentum states that the total momentum of an isolated system "
         "remains constant if no external net forces act upon it, as observed in elastic and inelastic collisions."),
        
        ("Chapter 5: Work-Energy Theorem",
         "The Work-Energy Theorem states that the total net work done by all forces acting on a particle "
         "equals the change in the particle's kinetic energy: W_net = Delta K. "
         "Work is defined as the dot product of force and displacement vectors: W = F * d * cos(theta)."),
        
        ("Chapter 6: Mechanical Energy and Gravitational Potential Energy",
         "Gravitational potential energy is the energy possessed by an object due to its vertical position in a gravitational field: "
         "U = m * g * h, where g is gravitational acceleration (9.8 m/s^2). "
         "In a conservative system without non-conservative friction, total mechanical energy E = K + U remains strictly constant."),
        
        ("Chapter 7: First Law of Thermodynamics",
         "The First Law of Thermodynamics is the law of conservation of energy applied to thermodynamic systems. "
         "It states that the change in internal energy Delta U of a closed system is equal to heat Q added to the system "
         "minus work W done by the system: Delta U = Q - W."),
        
        ("Chapter 8: Second Law of Thermodynamics and Entropy",
         "The Second Law of Thermodynamics states that the total entropy of an isolated system always increases over time "
         "for any spontaneous irreversible process. Heat cannot spontaneously flow from a colder body to a hotter body "
         "without external work input, setting the theoretical upper limit on heat engine efficiency."),
        
        ("Chapter 9: Coulomb's Law of Electrostatics",
         "Coulomb's Law quantifies the electrostatic force between two stationary electric point charges. "
         "The magnitude of the electrostatic force is directly proportional to the product of charges and inversely "
         "proportional to the square of the distance between them: F = k * |q1 * q2| / r^2, where k is Coulomb's constant."),
        
        ("Chapter 10: Snell's Law of Refraction",
         "Snell's Law of Refraction describes the relationship between the angles of incidence and refraction "
         "when light passes across the boundary between two isotropic media: n1 * sin(theta1) = n2 * sin(theta2), "
         "where n1 and n2 are the refractive indices of the respective optical media.")
    ]

    for title, text in pages_content:
        p = doc.new_page()
        rect = pymupdf.Rect(50, 50, 550, 750)
        p.insert_textbox(rect, f"{title}\n\n{text}")

    doc.save(file_path)
    doc.close()
    return file_path


def create_8_slide_pptx(file_path: str) -> str:
    """Create an 8-slide Electrical Machines Presentation."""
    prs = Presentation()
    slides_content = [
        ("Electric Motors Overview",
         "An electric motor is an electromechanical device that converts electrical energy into mechanical energy."),
        
        ("Motor Construction Components",
         "An electric motor consists of three primary components: the stationary stator, the rotating rotor, and the commutator."),
        
        ("Lorentz Force Principle",
         "Electric motors operate based on the Lorentz force law: F = q * (E + v x B). Current-carrying conductors in magnetic fields experience mechanical torque."),
        
        ("Direct Current (DC) Motors",
         "Direct current motors use Fleming's left-hand rule to determine the direction of force on armature conductors."),
        
        ("Alternating Current (AC) Induction Motors",
         "AC induction motors rely on a rotating magnetic field generated by three-phase stator windings to induce currents in the squirrel-cage rotor."),
        
        ("Transformer Working Principles",
         "Transformers transfer alternating electrical energy between circuits through mutual electromagnetic induction without changing frequency."),
        
        ("Transformer Turns Ratio Equation",
         "The turns ratio equation for ideal transformers relates primary and secondary voltages to winding turns: Vp / Vs = Np / Ns."),
        
        ("Power Inverters and Rectifier Circuits",
         "Rectifier circuits convert alternating current (AC) to direct current (DC), whereas inverters convert direct current (DC) to alternating current (AC).")
    ]

    for title, text in slides_content:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = text

    prs.save(file_path)
    return file_path


def create_8_topic_mp4(file_path: str) -> str:
    """Create an MP4 lecture video containing 8 distinct fluid mechanics concepts."""
    wav_path = file_path.replace(".mp4", "_speech.wav")
    spoken_text = (
        "According to Bernoulli principle, when fluid velocity increases, static pressure decreases. "
        "Fluid flow regime depends on Reynolds number, transitioning from smooth laminar flow to turbulent flow. "
        "The continuity equation for incompressible fluids states that cross-sectional area times velocity remains constant. "
        "The Venturi effect describes fluid pressure reduction when a fluid flows through a constricted narrow pipe section. "
        "Fluid dynamic viscosity generates internal shear resistance, governed by Poiseuille flow in cylindrical conduits. "
        "Aerodynamic lift is generated when an asymmetric cambered airfoil creates lower static pressure on its upper surface. "
        "Boundary layer separation occurs at excessive angles of attack, causing aerodynamic stall and loss of lift. "
        "Pascal principle states that pressure applied to an enclosed static fluid is transmitted undiminished throughout."
    )

    try:
        import win32com.client
        speaker = win32com.client.Dispatch("SAPI.SpVoice")
        stream = win32com.client.Dispatch("SAPI.SpFileStream")
        stream.Open(wav_path, 3, False)
        speaker.AudioOutputStream = stream
        speaker.Speak(spoken_text)
        stream.Close()
    except Exception as e:
        print(f"Warning: SAPI voice generation failed: {e}")
        wav_path = None

    ffmpeg_exe = get_ffmpeg_exe()
    if ffmpeg_exe and wav_path and os.path.exists(wav_path):
        cmd = [
            ffmpeg_exe, "-y",
            "-f", "lavfi", "-i", "color=c=darkblue:s=320x240:r=1",
            "-i", wav_path,
            "-c:v", "libx264", "-pix_fmt", "yuv420p",
            "-c:a", "aac", "-shortest", file_path
        ]
        subprocess.run(cmd, capture_output=True, text=True)

    if wav_path and os.path.exists(wav_path):
        try:
            os.remove(wav_path)
        except Exception:
            pass

    return file_path


def setup_phase5b_benchmark():
    print("=" * 80)
    print("SETUP PHASE 5B: COMPREHENSIVE 40-CASE MULTIMODAL BENCHMARK")
    print("=" * 80)

    client = httpx.Client(base_url=BASE_URL, timeout=120.0)

    # 1. Create Target Course 1
    print("\n[Step 1] Creating Target Benchmark Course...")
    c1_res = client.post("/api/courses", json={
        "title": "Phase 5B Multimodal Benchmark Course",
        "description": "Comprehensive benchmark course covering mechanics, electrical machines, fluids, and neural networks."
    })
    assert c1_res.status_code == 201
    course_1_id = c1_res.json()["id"]
    print(f"--> Target Course created: ID={course_1_id}")

    # 2. Create Isolated Course 2
    print("\n[Step 2] Creating Isolated Course 2...")
    c2_res = client.post("/api/courses", json={
        "title": "Phase 5B World History (Isolated Course)",
        "description": "Isolated course with zero physics or engineering material."
    })
    assert c2_res.status_code == 201
    course_2_id = c2_res.json()["id"]
    print(f"--> Isolated Course created: ID={course_2_id}")

    # 3. Create & Upload 10-page PDF
    pdf_path = os.path.join(os.path.dirname(__file__), "benchmark_physics_guide.pdf")
    create_10_page_pdf(pdf_path)
    print("\n[Step 3] Uploading 10-page Physics Guide PDF...")
    with open(pdf_path, "rb") as f:
        pdf_res = client.post(f"/api/courses/{course_1_id}/upload", files={"file": ("benchmark_physics_guide.pdf", f, "application/pdf")})
    assert pdf_res.status_code == 201
    pdf_sid = pdf_res.json()["source_id"]
    print(f"--> PDF uploaded: source_id={pdf_sid}")

    # 4. Create & Upload 8-slide PPTX
    pptx_path = os.path.join(os.path.dirname(__file__), "benchmark_electrical_machines.pptx")
    create_8_slide_pptx(pptx_path)
    print("\n[Step 4] Uploading 8-slide Electrical Machines PPTX...")
    with open(pptx_path, "rb") as f:
        pptx_res = client.post(f"/api/courses/{course_1_id}/upload", files={"file": ("benchmark_electrical_machines.pptx", f, "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert pptx_res.status_code == 201
    pptx_sid = pptx_res.json()["source_id"]
    print(f"--> PPTX uploaded: source_id={pptx_sid}")

    # 5. Create & Upload 8-topic MP4 Video
    mp4_path = os.path.join(os.path.dirname(__file__), "benchmark_fluid_mechanics.mp4")
    create_8_topic_mp4(mp4_path)
    print("\n[Step 5] Uploading 8-topic Fluid Mechanics MP4...")
    with open(mp4_path, "rb") as f:
        mp4_res = client.post(f"/api/courses/{course_1_id}/upload", files={"file": ("benchmark_fluid_mechanics.mp4", f, "video/mp4")})
    assert mp4_res.status_code == 201
    mp4_sid = mp4_res.json()["source_id"]
    print(f"--> MP4 uploaded: source_id={mp4_sid}")

    # 6. Add YouTube video
    print(f"\n[Step 6] Adding YouTube video source ({REAL_YOUTUBE_URL})...")
    yt_res = client.post(f"/api/courses/{course_1_id}/youtube", json={"url": REAL_YOUTUBE_URL})
    if yt_res.status_code == 201:
        yt_sid = yt_res.json()["source_id"]
        print(f"--> YouTube source added: source_id={yt_sid}")
    else:
        # Fallback to existing YouTube source ID in DB if already ingested
        from app.core.database import SessionLocal
        from app.models.source import Source
        db_s = SessionLocal()
        existing_yt = db_s.query(Source).filter(Source.source_type == "youtube", Source.processing_status == "completed").first()
        yt_sid = existing_yt.id if existing_yt else None
        db_s.close()
        print(f"--> YouTube online ingestion status: {yt_res.status_code}. Using fallback source: {yt_sid}")

    # Clean up local generated temp files
    for p in [pdf_path, pptx_path, mp4_path]:
        if os.path.exists(p):
            try:
                os.remove(p)
            except Exception:
                pass

    # 7. Construct 40-case Benchmark Dataset
    print("\n[Step 7] Constructing 40-case verified evaluation benchmark dataset...")
    cases = []

    # A. 10 PDF Test Cases (1 per page)
    pdf_cases = [
        ("eval-pdf-01", 1, "What is Newton's first law of motion and what property of matter resists velocity changes?", "Inertia"),
        ("eval-pdf-02", 2, "What is Newton's second law of motion and its mathematical equation?", "F = m * a"),
        ("eval-pdf-03", 3, "What does Newton's third law of motion state regarding action and reaction?", "Equal in magnitude and opposite in direction"),
        ("eval-pdf-04", 4, "What is linear momentum and when is the total momentum of an isolated system conserved?", "p = m * v"),
        ("eval-pdf-05", 5, "What does the work-energy theorem state about the work done on a particle?", "W_net = Delta K"),
        ("eval-pdf-06", 6, "What is gravitational potential energy and what is its formula?", "U = m * g * h"),
        ("eval-pdf-07", 7, "What is the first law of thermodynamics equation relating internal energy, heat, and work?", "Delta U = Q - W"),
        ("eval-pdf-08", 8, "What does the second law of thermodynamics state about the entropy of an isolated system?", "Total entropy always increases over time"),
        ("eval-pdf-09", 9, "What is Coulomb's law of electrostatics and how does force vary with distance?", "F = k * |q1 * q2| / r^2"),
        ("eval-pdf-10", 10, "What is Snell's law of refraction relating angles and refractive indices?", "n1 * sin(theta1) = n2 * sin(theta2)")
    ]
    for cid, page, q, ans in pdf_cases:
        cases.append({
            "id": cid,
            "course_id": course_1_id,
            "question": q,
            "expected_answer": ans,
            "expected_grounded": True,
            "expected_sources": [{"source_id": pdf_sid, "file_name": "benchmark_physics_guide.pdf", "page": page, "modality": "pdf"}]
        })

    # B. 8 PPTX Test Cases (1 per slide)
    pptx_cases = [
        ("eval-pptx-01", 1, "What is an electric motor according to the presentation slides?", "Electromechanical device converting electrical to mechanical energy"),
        ("eval-pptx-02", 2, "What are the three primary components of an electric motor mentioned in the slides?", "Stator, rotor, and commutator"),
        ("eval-pptx-03", 3, "What is the Lorentz force law formula that governs electric motor operation?", "F = q * (E + v x B)"),
        ("eval-pptx-04", 4, "Which rule is used to determine the force direction on armature conductors in DC motors?", "Fleming's left-hand rule"),
        ("eval-pptx-05", 5, "What type of magnetic field do AC induction motors rely on to induce currents in the rotor?", "Rotating magnetic field"),
        ("eval-pptx-06", 6, "Through what principle do transformers transfer electrical energy between circuits?", "Mutual electromagnetic induction"),
        ("eval-pptx-07", 7, "What is the turns ratio equation relating voltages and turns for an ideal transformer?", "Vp / Vs = Np / Ns"),
        ("eval-pptx-08", 8, "What is the difference in function between rectifier circuits and power inverters?", "Rectifiers convert AC to DC, inverters convert DC to AC")
    ]
    for cid, slide, q, ans in pptx_cases:
        cases.append({
            "id": cid,
            "course_id": course_1_id,
            "question": q,
            "expected_answer": ans,
            "expected_grounded": True,
            "expected_sources": [{"source_id": pptx_sid, "file_name": "benchmark_electrical_machines.pptx", "slide": slide, "modality": "pptx"}]
        })

    # C. 8 MP4 Video Test Cases (Timestamps across the lecture)
    mp4_cases = [
        ("eval-mp4-01", 0.0, 15.0, "According to Bernoulli's principle in the lecture video, what happens to static pressure when velocity increases?", "Static pressure decreases"),
        ("eval-mp4-02", 5.0, 20.0, "What dimensionless parameter determines the transition between laminar and turbulent flow?", "Reynolds number"),
        ("eval-mp4-03", 12.0, 30.0, "What does the continuity equation state for incompressible fluids flowing through a conduit?", "Area times velocity remains constant"),
        ("eval-mp4-04", 20.0, 35.0, "What is the Venturi effect described in the fluid mechanics lecture video?", "Fluid pressure reduction when flowing through a constricted narrow section"),
        ("eval-mp4-05", 28.0, 45.0, "What property generates internal shear resistance in Poiseuille pipe flow?", "Fluid dynamic viscosity"),
        ("eval-mp4-06", 36.0, 52.0, "How is aerodynamic lift generated on an asymmetric cambered airfoil?", "Lower static pressure on its upper surface"),
        ("eval-mp4-07", 44.0, 60.0, "What phenomenon occurs at excessive angles of attack leading to aerodynamic stall?", "Boundary layer separation"),
        ("eval-mp4-08", 50.0, 65.0, "What does Pascal's principle state about pressure applied to an enclosed static fluid?", "Pressure is transmitted undiminished in all directions")
    ]
    for cid, t0, t1, q, ans in mp4_cases:
        cases.append({
            "id": cid,
            "course_id": course_1_id,
            "question": q,
            "expected_answer": ans,
            "expected_grounded": True,
            "expected_sources": [{"source_id": mp4_sid, "file_name": "benchmark_fluid_mechanics.mp4", "start_seconds": t0, "end_seconds": t1, "modality": "video"}]
        })

    # D. 4 YouTube Test Cases
    if yt_sid:
        yt_cases = [
            ("eval-yt-01", 180.0, 240.0, "How many neurons make up the input layer for a 28 by 28 pixel image in the neural network video?", "784 neurons"),
            ("eval-yt-02", 180.0, 240.0, "What range of numerical values can the activation of a neuron take?", "Between 0 and 1"),
            ("eval-yt-03", 190.0, 265.0, "How many output neurons are in the final layer for recognizing handwritten digits?", "10 output neurons"),
            ("eval-yt-04", 440.0, 560.0, "What subcomponents do hidden layers detect in the neural network video?", "Patterns of edges and loops")
        ]
        for cid, t0, t1, q, ans in yt_cases:
            cases.append({
                "id": cid,
                "course_id": course_1_id,
                "question": q,
                "expected_answer": ans,
                "expected_grounded": True,
                "expected_sources": [{"source_id": yt_sid, "start_seconds": t0, "end_seconds": t1, "modality": "youtube"}]
            })

    # E. 5 Unsupported Off-Material Test Cases
    unsupported_cases = [
        ("eval-unsupported-01", "What is the capital of ancient Mesopotamia and what were the complete laws of Hammurabi?"),
        ("eval-unsupported-02", "What are the official tournament rules and scoring of professional cricket including the LBW rule?"),
        ("eval-unsupported-03", "Who were the primary emperors of the Ming Dynasty in imperial Chinese history?"),
        ("eval-unsupported-04", "Explain the detailed biological mechanism of chlorophyll synthesis in marine cyanobacteria."),
        ("eval-unsupported-05", "What are the exact FIFA regulations defining an offside violation in professional soccer?")
    ]
    for cid, q in unsupported_cases:
        cases.append({
            "id": cid,
            "course_id": course_1_id,
            "question": q,
            "expected_answer": None,
            "expected_grounded": False,
            "expected_sources": []
        })

    # F. 5 Course-Isolation Leakage Test Cases (Submitted to Course 2 asking for Course 1's content)
    isolation_cases = [
        ("eval-isolation-01", "What is Newton's second law and the formula from the physics document?"),
        ("eval-isolation-02", "What are the three components of an electric motor mentioned in the slides?"),
        ("eval-isolation-03", "According to Bernoulli's principle in the video, what happens to fluid pressure?"),
        ("eval-isolation-04", "How many neurons make up the input layer of the neural network?"),
        ("eval-isolation-05", "What is the continuity equation for incompressible fluid flow?")
    ]
    for cid, q in isolation_cases:
        cases.append({
            "id": cid,
            "course_id": course_2_id,  # Query submitted against empty Course 2
            "question": q,
            "expected_answer": None,
            "expected_grounded": False,
            "expected_sources": [],
            "is_course_isolation": True,
            "source_course_id": course_1_id
        })

    # Save benchmark dataset to JSON file
    dataset_file = os.path.join(os.path.dirname(__file__), "..", "app", "services", "evaluation", "benchmark_dataset_phase5b.json")
    with open(dataset_file, "w", encoding="utf-8") as f:
        json.dump(cases, f, indent=2)

    print(f"\n--> Benchmark Dataset created successfully with {len(cases)} test cases!")
    print(f"--> Saved to: {dataset_file}")
    return course_1_id, course_2_id, dataset_file, len(cases)


if __name__ == "__main__":
    setup_phase5b_benchmark()
