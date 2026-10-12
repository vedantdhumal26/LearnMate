import io
import pytest
from unittest.mock import patch, MagicMock
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

from app.models.course import Course
from app.models.source import Source
from app.models.document_chunk import DocumentChunk
from app.services.vector_store.qdrant_store import get_vector_store
from app.services.rag.tutor_service import TutorService
from tests.conftest import MockLLMService

def generate_synthetic_pptx() -> bytes:
    """
    Generate a 3-slide synthetic PPTX in memory:
    - Slide 1: Title and introduction
    - Slide 2: Physics formula (Delta U = Q - W) and details
    - Slide 3: Embedded diagram / image and description
    """
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Normal text & title
    slide1 = prs.slides.add_slide(blank_layout)
    tx_box1 = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
    tf1 = tx_box1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Thermodynamics and Kinetic Theory"
    p1.font.size = Pt(28)
    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Fundamental concepts of heat, internal energy, and thermodynamic systems."

    # Slide 2: Formula & specific details
    slide2 = prs.slides.add_slide(blank_layout)
    tx_box2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
    tf2 = tx_box2.text_frame
    p2_title = tf2.paragraphs[0]
    p2_title.text = "First Law of Thermodynamics"
    p2_title.font.size = Pt(24)
    p2_formula = tf2.add_paragraph()
    p2_formula.text = "The change in internal energy is given by: Delta U = Q - W"
    p2_desc = tf2.add_paragraph()
    p2_desc.text = "Here Delta U represents internal energy change, Q represents heat added, and W is work done."

    # Slide 3: Embedded Image / Diagram
    slide3 = prs.slides.add_slide(blank_layout)
    tx_box3 = slide3.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tf3 = tx_box3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = "Carnot Cycle Pressure-Volume Diagram"

    # Create a real small PNG image in-memory
    img = Image.new("RGB", (200, 200), color=(50, 120, 200))
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    img_byte_arr.seek(0)
    slide3.shapes.add_picture(img_byte_arr, Inches(1), Inches(1.5), Inches(4), Inches(3))

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()

@pytest.fixture
def sample_pptx() -> bytes:
    return generate_synthetic_pptx()

def test_pptx_upload_success(client, sample_pptx):
    """1. Test successful PPTX upload and 201 response."""
    course_res = client.post("/api/courses", json={"title": "Thermodynamics 101"})
    assert course_res.status_code == 201
    course_id = course_res.json()["id"]

    files = {"file": ("thermo_lecture.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    response = client.post(f"/api/courses/{course_id}/upload", files=files)
    assert response.status_code == 201
    data = response.json()
    assert "source_id" in data
    assert data["file_name"] == "thermo_lecture.pptx"
    assert data["status"] == "completed"

def test_pptx_slide_text_extraction(client, db_session, sample_pptx):
    """2. Test that text is extracted slide-by-slide accurately."""
    course_res = client.post("/api/courses", json={"title": "Thermal Physics"})
    course_id = course_res.json()["id"]

    files = {"file": ("thermo_lecture.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    upload_res = client.post(f"/api/courses/{course_id}/upload", files=files)
    source_id = upload_res.json()["source_id"]

    chunks = db_session.query(DocumentChunk).filter(DocumentChunk.source_id == source_id).all()
    assert len(chunks) >= 3

    # Check Slide 1 content
    slide1_chunks = [c for c in chunks if c.slide_number == 1]
    assert any("Thermodynamics" in c.content for c in slide1_chunks)

    # Check Slide 2 formula content
    slide2_chunks = [c for c in chunks if c.slide_number == 2]
    assert any("Delta U = Q - W" in c.content for c in slide2_chunks)

def test_pptx_slide_number_preserved(client, db_session, sample_pptx):
    """3. Test that 1-indexed slide_number is strictly preserved and page_number is None."""
    course_res = client.post("/api/courses", json={"title": "Slide Numbering Test"})
    course_id = course_res.json()["id"]

    files = {"file": ("slides.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    upload_res = client.post(f"/api/courses/{course_id}/upload", files=files)
    source_id = upload_res.json()["source_id"]

    chunks = db_session.query(DocumentChunk).filter(DocumentChunk.source_id == source_id).all()
    slide_numbers = {c.slide_number for c in chunks}
    assert 1 in slide_numbers
    assert 2 in slide_numbers
    assert 3 in slide_numbers

    for c in chunks:
        assert c.page_number is None
        assert c.slide_number is not None

def test_pptx_chunks_created(client, db_session, sample_pptx):
    """4. Test that chunks have content_hash, course_id, and modality."""
    course_res = client.post("/api/courses", json={"title": "Chunk Creation Test"})
    course_id = course_res.json()["id"]

    files = {"file": ("chunks.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    upload_res = client.post(f"/api/courses/{course_id}/upload", files=files)
    source_id = upload_res.json()["source_id"]

    chunks = db_session.query(DocumentChunk).filter(DocumentChunk.source_id == source_id).all()
    assert len(chunks) >= 3
    for chunk in chunks:
        assert chunk.course_id == course_id
        assert len(chunk.content_hash) == 64
        assert chunk.modality in ["pptx_text", "pptx_visual"]

def test_pptx_vectors_indexed(client, sample_pptx):
    """5. Test that PPTX vectors are indexed in Qdrant with slide metadata."""
    course_res = client.post("/api/courses", json={"title": "Vector Index Test"})
    course_id = course_res.json()["id"]

    files = {"file": ("indexed.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    client.post(f"/api/courses/{course_id}/upload", files=files)

    store = get_vector_store()
    results = store.search_vectors(
        query_vector=[0.0] * 384,
        course_id=course_id,
        limit=10,
        score_threshold=0.0
    )
    assert len(results) >= 3
    slide_numbers = {r.slide_number for r in results if r.slide_number is not None}
    assert len(slide_numbers) >= 2
    for r in results:
        assert r.file_name == "indexed.pptx"
        assert r.modality in ["pptx_text", "pptx_visual"]

def test_pptx_tutor_grounded_answer(client, sample_pptx):
    """6. Test tutor provides grounded answer from PPTX content."""
    course_res = client.post("/api/courses", json={"title": "Tutor Test Course"})
    course_id = course_res.json()["id"]

    files = {"file": ("thermo.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    client.post(f"/api/courses/{course_id}/upload", files=files)

    # Use mock LLM
    mock_llm = MockLLMService(response_text="The formula for internal energy is Delta U = Q - W according to thermo.pptx — Slide 2.")
    with patch("app.services.rag.tutor_service.get_llm_service", return_value=mock_llm):
        tutor_res = client.post("/api/tutor/chat", json={
            "course_id": course_id,
            "message": "What is the formula for the change in internal energy?"
        })
        assert tutor_res.status_code == 200
        data = tutor_res.json()
        assert data["grounded"] is True
        assert "Delta U = Q - W" in data["answer"]
        assert len(data["sources"]) > 0

def test_pptx_tutor_returns_slide_citation(client, sample_pptx):
    """7. Test tutor returns slide citation (slide=2, page=null)."""
    course_res = client.post("/api/courses", json={"title": "Citation Test Course"})
    course_id = course_res.json()["id"]

    files = {"file": ("thermodynamics.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    client.post(f"/api/courses/{course_id}/upload", files=files)

    mock_llm = MockLLMService(response_text="Based on thermodynamics.pptx — Slide 2, Delta U = Q - W.")
    with patch("app.services.rag.tutor_service.get_llm_service", return_value=mock_llm):
        tutor_res = client.post("/api/tutor/chat", json={
            "course_id": course_id,
            "message": "internal energy formula Delta U = Q - W"
        })
        assert tutor_res.status_code == 200
        data = tutor_res.json()
        assert data["grounded"] is True
        assert len(data["sources"]) > 0
        citation = data["sources"][0]
        assert citation["file_name"] == "thermodynamics.pptx"
        assert citation["slide"] == 2
        assert citation["page"] is None

def test_pdf_behavior_still_works(client, sample_physics_pdf):
    """8. Test that PDF citations still return page number and slide=null."""
    course_res = client.post("/api/courses", json={"title": "PDF Regression Course"})
    course_id = course_res.json()["id"]

    files = {"file": ("physics.pdf", io.BytesIO(sample_physics_pdf), "application/pdf")}
    client.post(f"/api/courses/{course_id}/upload", files=files)

    mock_llm = MockLLMService(response_text="According to Newton's Second Law on physics.pdf — Page 2, F = ma.")
    with patch("app.services.rag.tutor_service.get_llm_service", return_value=mock_llm):
        tutor_res = client.post("/api/tutor/chat", json={
            "course_id": course_id,
            "message": "What is Newton's second law F = m * a?"
        })
        assert tutor_res.status_code == 200
        data = tutor_res.json()
        assert data["grounded"] is True
        assert len(data["sources"]) > 0
        citation = data["sources"][0]
        assert citation["file_name"] == "physics.pdf"
        assert citation["page"] is not None
        assert citation["slide"] is None

def test_course_isolation_with_pptx(client, sample_pptx):
    """9. Test that Course 1 PPTX content is isolated from Course 2."""
    c1_res = client.post("/api/courses", json={"title": "Course 1 PPTX"})
    c1_id = c1_res.json()["id"]
    c2_res = client.post("/api/courses", json={"title": "Course 2 Empty"})
    c2_id = c2_res.json()["id"]

    files = {"file": ("thermo.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    client.post(f"/api/courses/{c1_id}/upload", files=files)

    # Ask in Course 2 (which has no materials)
    tutor_res = client.post("/api/tutor/chat", json={
        "course_id": c2_id,
        "message": "What is the formula for the change in internal energy Delta U?"
    })
    assert tutor_res.status_code == 200
    data = tutor_res.json()
    assert data["grounded"] is False
    assert data["sources"] == []

def test_unsupported_ppt_extension_rejected(client):
    """10. Test that .ppt binary files are explicitly rejected with Phase 4A message."""
    course_res = client.post("/api/courses", json={"title": "PPT Rejection Test"})
    course_id = course_res.json()["id"]

    files = {"file": ("legacy_lecture.ppt", io.BytesIO(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"), "application/vnd.ms-powerpoint")}
    response = client.post(f"/api/courses/{course_id}/upload", files=files)
    assert response.status_code == 400
    assert response.json()["detail"] == "Only .pptx slide decks are supported in Phase 4A."

def test_duplicate_pptx_handling(client, sample_pptx):
    """11. Test duplicate PPTX upload returns 409 Conflict."""
    course_res = client.post("/api/courses", json={"title": "Duplicate PPTX Test"})
    course_id = course_res.json()["id"]

    files1 = {"file": ("duplicate.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res1 = client.post(f"/api/courses/{course_id}/upload", files=files1)
    assert res1.status_code == 201

    files2 = {"file": ("duplicate.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res2 = client.post(f"/api/courses/{course_id}/upload", files=files2)
    assert res2.status_code == 409
    assert "already been uploaded" in res2.json()["detail"]

def test_visual_extraction_does_not_break_ingestion(client, db_session, sample_pptx):
    """12. Test that visual extraction succeeds, attaches visual_description when available, and fails gracefully."""
    course_res = client.post("/api/courses", json={"title": "Visual Extraction Test"})
    course_id = course_res.json()["id"]

    # Test with vision mock returning a diagram description
    mock_desc = "Diagram showing Carnot cycle pressure vs volume with adiabatic and isothermal curves."
    with patch("app.services.llm.vision_service.VisionService.describe_image", return_value=mock_desc):
        files = {"file": ("visuals.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
        res = client.post(f"/api/courses/{course_id}/upload", files=files)
        assert res.status_code == 201
        source_id = res.json()["source_id"]

        chunks = db_session.query(DocumentChunk).filter(DocumentChunk.source_id == source_id).all()
        visual_chunks = [c for c in chunks if c.modality == "pptx_visual"]
        assert len(visual_chunks) >= 1
        assert visual_chunks[0].visual_description == mock_desc

    # Test that when vision throws an exception, ingestion continues and does NOT crash
    course_res2 = client.post("/api/courses", json={"title": "Vision Fallback Test"})
    course_id2 = course_res2.json()["id"]

    with patch("app.services.llm.vision_service.VisionService.describe_image", side_effect=Exception("Vision API rate limit")):
        files2 = {"file": ("visuals_fallback.pptx", io.BytesIO(sample_pptx), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
        res2 = client.post(f"/api/courses/{course_id2}/upload", files=files2)
        assert res2.status_code == 201
        data2 = res2.json()
        assert data2["status"] == "completed"

        chunks2 = db_session.query(DocumentChunk).filter(DocumentChunk.source_id == data2["source_id"]).all()
        # Slide text chunks must still exist
        assert len(chunks2) >= 2
        slide2_chunks = [c for c in chunks2 if c.slide_number == 2]
        assert len(slide2_chunks) >= 1
